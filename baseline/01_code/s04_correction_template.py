"""
AGEVAL Step 4 — Correction Template Generator
Run: python baseline/01_code/s04_correction_template.py --data baseline/04_data/dataRaw/test_data_honduras_n2052.xlsx

Reuses the same quality rules and check functions as s03_quality_check.py
(imported, not duplicated) to identify flagged records, then builds one Excel
worksheet per data sheet — general/main sheet plus every loop/repeat sheet —
for the survey supervision team to review and correct.

For each ORIGINAL variable that has at least one flagged issue, four columns
are generated on its own sheet's tab:
    <var> - Original value
    <var> - Flag reason
    <var> - Corrected value   (blank, for supervisors to fill in)
    <var> - Correction reason (blank, for supervisors to fill in)

A variable's columns are only populated for a given record when that record
actually has an issue on that variable; otherwise the whole group of 4
columns is left blank for that row.

Calculated variables (`surv_type == "calculate"`, e.g. `area_total_ha`,
`farm_area_total_ha`) are used to DETECT problems but are never shown
themselves. Instead, a flagged calculated variable is decomposed — via
`surv_calculation`, recursively through any intermediate calculated
variables — into the underlying original variable(s) that should actually be
corrected:
  - If those original variables live on the SAME row as the flagged
    calculated value (e.g. a plot-level `area_total_ha` outlier -> that
    plot's own `area_total`/`area_unit`), they're added to that same row,
    tagged as a plot/observation-level reason.
  - If the flagged calculated value is an AGGREGATE computed from another
    sheet (e.g. a respondent-level `farm_area_total_ha` outlier -> every one
    of that respondent's plot-level `area_total`/`area_unit` observations),
    every underlying observation for that respondent is pulled in, tagged as
    an aggregate/respondent-level reason.
A single observation can end up with both a plot-level and an
aggregate-level reason (or just one) — both are shown, joined together, in
that variable's "Flag reason" column.

This logic is fully generic: it relies only on `surv_type`/`surv_calculation`
relationships in the dictionary, never on hard-coded variable names.
"""

import re
import pandas as pd
import numpy as np
import os
import sys
import yaml
import argparse
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import s03_quality_check as qc

# ── Reuse config, dictionary paths, and quality-check functions from s03 ───────
# Read config
CFG = os.path.join(os.path.dirname(__file__), "s00_config.yaml")
with open(CFG, "r") as f:
    config = yaml.safe_load(f)
SURVEY_ROUND = config["project"]["survey_round"]

# Paths
BASE       = os.path.normpath(os.path.join(os.path.dirname(CFG), config["paths"]["base"]))
MASTER     = os.path.join(BASE, config["paths"]["dictionary_master"])
DICT       = os.path.join(BASE, config["paths"]["dictionary_personalized"].format(survey_round=SURVEY_ROUND))
CORRECTIONS = os.path.join(BASE, config["paths"]["correction_files"].format(survey_round=SURVEY_ROUND))


# Columns always shown to identify a record, regardless of which variable
# triggered the issue. Only the ones actually present in a given sheet are
# used (see `_id_cols_for`), so this list can be a superset.
ID_COLS      = ["respondent_id", "start", "end", "surveyDate", "deviceid"]
ID_COLS_LOOP = ["respondent_id", "_obs_id", "surveyDate", "start", "end", "deviceid"]

# ── Styling ──────────────────────────────────────────────────────────────────
FONT_NAME = "Arial"

TITLE_FONT   = Font(name=FONT_NAME, size=13, bold=True, color="1D4ED8")
SUBTITLE_FONT = Font(name=FONT_NAME, size=9, italic=True, color="64748B")
GROUP_FONT   = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
HEADER_FONT  = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
ID_HEADER_FONT = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
BODY_FONT    = Font(name=FONT_NAME, size=10)
INPUT_FONT   = Font(name=FONT_NAME, size=10, color="1A1A1A")

GROUP_FILL   = PatternFill("solid", fgColor="1D4ED8")
ID_FILL      = PatternFill("solid", fgColor="1E293B")
INPUT_FILL   = PatternFill("solid", fgColor="FEF3C7")   # yellow — supervisor input
STRIPE_FILL  = PatternFill("solid", fgColor="F8FAFC")   # light row stripe

THIN  = Side(style="thin", color="CBD5E1")
THICK = Side(style="medium", color="94A3B8")

CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP   = Alignment(horizontal="left", vertical="center", wrap_text=True)

GROUP_HEADER_ROW  = 3
COL_HEADER_ROW    = 4
DATA_START_ROW    = 5
SUBCOLS = ["Original value", "Flag reason", "Corrected value", "Correction reason"]

ID_DISPLAY_NAMES = {
    "respondent_id": "Respondent ID",
    "start": "Start",
    "end": "End",
    "surveyDate": "Survey date",
    "deviceid": "Device ID",
    "_obs_id": "Observation ID",
}
ID_WIDTHS = {
    "respondent_id": 16, "start": 20, "end": 20, "surveyDate": 14,
    "deviceid": 16, "_obs_id": 14,
}


# ── Recompute quality-check results, per sheet (mirrors run_quality_check's
#    core loop, without the HTML rendering step) ────────────────────────────
def compute_quality_results_by_sheet(data_path, date_start=None, date_end=None, date_col="surveyDate", dict_path=None):
    """
    Returns a dict:
      dict_df:         full dictionary DataFrame (all variables, not just
                        quality_include==1 — needed to resolve calculated ->
                        original variable relationships)
      main_sheet_name: name of the general/main sheet
      id_col:          respondent-level id column name (e.g. "respondent_id")
      sheets:          {sheet_name: report_data} — date-filtered, reset-index
                        DataFrame for each sheet (general + every loop)
      sheet_id_field:  {sheet_name: field name used to key issues/rows for
                        that sheet} — `id_col` for the general sheet,
                        "_report_id" for loop sheets
      results:         {sheet_name: {variable_name: {label, checks,
                        flagged_ids}}}
      duplicates:      duplicate-check results (general sheet only)
    """
    dict_df = pd.read_csv(dict_path or DICT)
    qc_vars = dict_df[dict_df["quality_include"] == 1]

    sheets_raw, main_sheet_name = qc.load_sheets(data_path, qc.config)
    main_data = sheets_raw[main_sheet_name]

    main_reference = qc.filter_by_date(main_data, None, date_end, date_col)
    main_report    = qc.filter_by_date(main_data, date_start, date_end, date_col).reset_index(drop=True)

    id_col = "respondent_id" if "respondent_id" in main_report.columns else main_report.columns[0]

    # Calculated variables are checked directly on their own column (their
    # variable_name IS the actual data column) — no separate output column
    # concept anymore, so this is now a straight intersection.
    compare_cols = [v for v in qc_vars["variable_name"] if v in main_report.columns]

    dup_checks, dup_flagged = qc.check_duplicates(main_report, id_col, compare_cols, config["quality"]["duplicate_pct"])

    general_results, _, _ = qc.run_variable_checks(main_report, main_reference, qc_vars, id_col)

    sheets_report  = {main_sheet_name: main_report}
    sheets_results = {main_sheet_name: general_results}
    sheet_id_field = {main_sheet_name: id_col}

    for sheet_name, raw_loop_df in sheets_raw.items():
        if sheet_name == main_sheet_name:
            continue
        linked = qc.link_loop_to_main(raw_loop_df, main_data, id_col=id_col, date_col=date_col)
        loop_reference = qc.filter_by_date(linked, None, date_end, date_col)
        loop_report    = qc.filter_by_date(linked, date_start, date_end, date_col).reset_index(drop=True)

        loop_results, _, _ = qc.run_variable_checks(loop_report, loop_reference, qc_vars, "_report_id")

        sheets_report[sheet_name]  = loop_report
        sheets_results[sheet_name] = loop_results
        sheet_id_field[sheet_name] = "_report_id"

    return {
        "dict_df": dict_df,
        "main_sheet_name": main_sheet_name,
        "id_col": id_col,
        "sheets": sheets_report,
        "sheet_id_field": sheet_id_field,
        "results": sheets_results,
        "duplicates": {"label": "Duplicate detection", "checks": dup_checks, "flagged_ids": dup_flagged},
    }


# ── Calculated -> original variable resolution ──────────────────────────────
def _is_calculated(var_name, type_lookup):
    t = type_lookup.get(var_name)
    return isinstance(t, str) and t.strip().lower() == "calculate"


def resolve_leaf_vars(var_name, type_lookup, calc_lookup, _seen=None):
    """
    Recursively resolve a (possibly calculated) variable down to the set of
    original, non-calculated variables that feed it, following
    `surv_calculation` (e.g. "sum(${area_total_ha})" -> area_total_ha ->
    "if(${area_unit}=... ${area_total} ...)" -> {area_unit, area_total}).
    A variable not found in the dictionary at all is treated as original
    (there's nothing to decompose it into). Guards against circular
    references via `_seen`.
    """
    if _seen is None:
        _seen = set()
    if var_name in _seen:
        return set()
    _seen.add(var_name)

    if not _is_calculated(var_name, type_lookup):
        return {var_name}

    calc_expr = calc_lookup.get(var_name)
    deps = qc._extract_odk_vars(calc_expr) if pd.notna(calc_expr) else []
    if not deps:
        # Marked as calculated but no formula available to decompose it —
        # nothing usable to put in the template.
        return set()

    leaves = set()
    for dep in deps:
        leaves |= resolve_leaf_vars(dep, type_lookup, calc_lookup, _seen)
    return leaves


# ── Reshape flagged records into a per-sheet, per-ID, per-variable issue map ─
def add_issue(store, sheet_name, id_key, var_key, value, issue_text):
    store.setdefault(sheet_name, {}).setdefault(str(id_key), {}).setdefault(var_key, []).append(
        {"value": value, "issue": issue_text}
    )


def register_var(order_store, label_store, sheet_name, var_key, label):
    labels = label_store.setdefault(sheet_name, {})
    if var_key not in labels:
        order_store.setdefault(sheet_name, []).append(var_key)
        labels[var_key] = label


def _lookup_row(df, id_field, id_value):
    if id_field not in df.columns:
        return None
    matches = df[df[id_field].astype(str) == str(id_value)]
    return matches.iloc[0] if not matches.empty else None


def build_correction_maps(compute_result):
    """
    Returns:
      issues:     {sheet_name: {id_key(str): {var_key: [{"value","issue"}, ...]}}}
      var_order:  {sheet_name: [var_key, ...]} in first-seen order
      var_labels: {sheet_name: {var_key: display label}}
    Only original (non-calculated) variables ever appear as `var_key`.
    """
    dict_df        = compute_result["dict_df"]
    sheets         = compute_result["sheets"]
    sheet_id_field = compute_result["sheet_id_field"]
    main_sheet_name = compute_result["main_sheet_name"]
    id_col         = compute_result["id_col"]

    type_lookup  = dict_df.set_index("variable_name")["surv_type"].to_dict()
    calc_lookup  = dict_df.set_index("variable_name")["surv_calculation"].to_dict()
    label_lookup = (
        dict_df.set_index("variable_name")["label_spanish"].to_dict()
        if "label_spanish" in dict_df.columns else {}
    )

    # Which sheet each dictionary variable actually lives on (first match).
    var_to_sheet = {}
    for vname in dict_df["variable_name"].dropna().unique():
        for sheet_name, df in sheets.items():
            if vname in df.columns:
                var_to_sheet[vname] = sheet_name
                break

    # Pre-group each sheet's row positions by plain respondent id, so an
    # aggregate-level issue can pull in every observation for a respondent
    # without rescanning the sheet each time.
    respondent_groups = {
        sheet_name: (df.groupby(df[id_col].astype(str)).groups if id_col in df.columns else {})
        for sheet_name, df in sheets.items()
    }

    issues     = {}
    var_order  = {}
    var_labels = {}

    # ── Duplicates (general/main sheet only) ────────────────────────────────
    dup = compute_result["duplicates"]
    dup_flagged = dup.get("flagged_ids", [])
    for rec in dup_flagged:
        rid = rec["id"]
        if isinstance(rid, str) and " / " in rid and rec["issue"] == "Potential duplicate based on similarity":
            id_a, id_b = [x.strip() for x in rid.split(" / ", 1)]
            add_issue(issues, main_sheet_name, id_a, "duplicates", rec["value"], f"{rec['issue']} (paired with respondent_id {id_b})")
            add_issue(issues, main_sheet_name, id_b, "duplicates", rec["value"], f"{rec['issue']} (paired with respondent_id {id_a})")
        else:
            add_issue(issues, main_sheet_name, rid, "duplicates", rec["value"], rec["issue"])
    if dup_flagged:
        register_var(var_order, var_labels, main_sheet_name, "duplicates", dup.get("label", "Duplicate detection"))

    # ── Per-sheet variable results ──────────────────────────────────────────
    for sheet_name, results in compute_result["results"].items():
        id_field = sheet_id_field[sheet_name]
        sheet_df = sheets[sheet_name]

        for var_name, result in results.items():
            flagged = result.get("flagged_ids", [])
            if not flagged:
                continue
            label = result.get("label", var_name)

            if not _is_calculated(var_name, type_lookup):
                # Original variable — handled exactly as before.
                for rec in flagged:
                    add_issue(issues, sheet_name, rec["id"], var_name, rec["value"], rec["issue"])
                register_var(var_order, var_labels, sheet_name, var_name, label)
                continue

            # ── Calculated variable: never shown itself — decompose into
            #    the original variable(s) that should be corrected. ────────
            leaf_vars = resolve_leaf_vars(var_name, type_lookup, calc_lookup)
            if not leaf_vars:
                continue

            for rec in flagged:
                flagged_id = str(rec["id"])
                issue_text = rec["issue"]

                for leaf in sorted(leaf_vars):
                    leaf_sheet = var_to_sheet.get(leaf)
                    if leaf_sheet is None:
                        continue
                    leaf_label = label_lookup.get(leaf, leaf)

                    if leaf_sheet == sheet_name:
                        # Same-row case: e.g. a plot-level area_total_ha
                        # outlier -> that same plot's area_total/area_unit.
                        row = _lookup_row(sheet_df, id_field, flagged_id)
                        leaf_val = row.get(leaf) if row is not None else None
                        reason = (
                            f"Calculated variable '{var_name}' ({label}) flagged: {issue_text} "
                            f"— included at the plot/observation level"
                        )
                        add_issue(issues, sheet_name, flagged_id, leaf, leaf_val, reason)
                        register_var(var_order, var_labels, sheet_name, leaf, leaf_label)
                    else:
                        # Aggregate case: e.g. a respondent-level
                        # farm_area_total_ha outlier -> every one of that
                        # respondent's plot-level area_total/area_unit rows.
                        if id_field == id_col:
                            plain_rid = flagged_id
                        else:
                            row = _lookup_row(sheet_df, id_field, flagged_id)
                            plain_rid = str(row.get(id_col)) if row is not None else None
                        if plain_rid is None:
                            continue

                        leaf_sheet_df  = sheets[leaf_sheet]
                        leaf_id_field  = sheet_id_field[leaf_sheet]
                        row_positions  = respondent_groups.get(leaf_sheet, {}).get(plain_rid, [])
                        for pos in row_positions:
                            mrow = leaf_sheet_df.loc[pos]
                            row_id   = mrow.get(leaf_id_field, pos)
                            leaf_val = mrow.get(leaf)
                            reason = (
                                f"Aggregate calculated variable '{var_name}' ({label}) flagged for "
                                f"respondent {plain_rid}: {issue_text} — included at the aggregate/respondent level"
                            )
                            add_issue(issues, leaf_sheet, row_id, leaf, leaf_val, reason)
                            register_var(var_order, var_labels, leaf_sheet, leaf, leaf_label)

    return issues, var_order, var_labels


def _safe_sheet_title(name, used):
    """Excel worksheet titles: <=31 chars, no : \\ / ? * [ ], unique."""
    safe = re.sub(r'[:\\/?*\[\]]', "-", str(name))[:31]
    base, i = safe, 1
    while safe in used:
        suffix = f"~{i}"
        safe = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(safe)
    return safe


# ── Excel builder: one worksheet per data sheet ─────────────────────────────
def _write_correction_sheet(wb, title, id_cols, issues, var_order, var_labels, report_data, id_field, batch, sheet_label):
    ws = wb.create_sheet(title=title)
    n_id_cols = len(id_cols)
    n_flagged = len(issues)
    n_cols = n_id_cols + 4 * len(var_order)
    n_cols = max(n_cols, n_id_cols, 1)

    # ── Title / instructions rows ───────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = (
        f"AGEVAL Correction Template — Batch: {batch}   |   Sheet: {sheet_label}   |   "
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}   |   "
        f"Flagged records: {n_flagged}"
    )
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    sub_cell = ws.cell(row=2, column=1)
    sub_cell.value = (
        "Review the 'Original value' and 'Flag reason' columns for each flagged variable. "
        "Fill in the yellow 'Corrected value' and 'Correction reason' columns only where a "
        "correction is needed. Leave a variable's columns blank if it had no issue for that record."
    )
    sub_cell.font = SUBTITLE_FONT
    sub_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 28

    # ── ID column headers (merged vertically across the two header rows) ───
    for i, col_name in enumerate(id_cols, start=1):
        ws.merge_cells(start_row=GROUP_HEADER_ROW, start_column=i, end_row=COL_HEADER_ROW, end_column=i)
        cell = ws.cell(row=GROUP_HEADER_ROW, column=i)
        cell.value = ID_DISPLAY_NAMES.get(col_name, col_name)
        cell.font = ID_HEADER_FONT
        cell.fill = ID_FILL
        cell.alignment = CENTER_WRAP
        for r in (GROUP_HEADER_ROW, COL_HEADER_ROW):
            ws.cell(row=r, column=i).fill = ID_FILL
            ws.cell(row=r, column=i).border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    # ── Variable group headers ──────────────────────────────────────────────
    col_ptr = n_id_cols + 1
    group_col_start = {}
    for var_key in var_order:
        group_col_start[var_key] = col_ptr
        label = var_labels.get(var_key, var_key)
        group_title = label if var_key == "duplicates" else f"{var_key} — {label}"

        ws.merge_cells(start_row=GROUP_HEADER_ROW, start_column=col_ptr, end_row=GROUP_HEADER_ROW, end_column=col_ptr + 3)
        gcell = ws.cell(row=GROUP_HEADER_ROW, column=col_ptr)
        gcell.value = group_title
        gcell.font = GROUP_FONT
        gcell.fill = GROUP_FILL
        gcell.alignment = CENTER_WRAP

        for offset, subname in enumerate(SUBCOLS):
            c = col_ptr + offset
            hcell = ws.cell(row=COL_HEADER_ROW, column=c)
            hcell.value = subname
            hcell.font = HEADER_FONT
            hcell.fill = GROUP_FILL
            hcell.alignment = CENTER_WRAP
            is_left_edge = (offset == 0)
            hcell.border = Border(
                left=THICK if is_left_edge else THIN, right=THIN, top=THIN, bottom=THIN,
            )
            ws.cell(row=GROUP_HEADER_ROW, column=c).border = Border(
                left=THICK if is_left_edge else THIN, right=THIN, top=THIN, bottom=THIN,
            )

        col_ptr += 4

    ws.row_dimensions[GROUP_HEADER_ROW].height = 22
    ws.row_dimensions[COL_HEADER_ROW].height = 30

    # ── Data rows ────────────────────────────────────────────────────────────
    # Iterate physical records directly (not a dict keyed by ID) so that
    # genuine duplicate respondent_id records — exactly what the duplicate
    # check flags — each keep their own true start/end/deviceid values
    # instead of collapsing onto a single row.
    r_ptr = DATA_START_ROW
    for _, src_row in report_data.iterrows():
        id_key = str(src_row.get(id_field))
        if id_key not in issues:
            continue
        stripe = ((r_ptr - DATA_START_ROW) % 2 == 1)

        # Identifying columns
        for i, col_name in enumerate(id_cols, start=1):
            cell = ws.cell(row=r_ptr, column=i)
            val = src_row.get(col_name) if col_name in src_row else None
            cell.value = "" if pd.isna(val) else val
            cell.font = BODY_FONT
            cell.alignment = LEFT_WRAP
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if col_name == "surveyDate":
                cell.number_format = "yyyy-mm-dd"
            elif col_name in ("start", "end"):
                cell.number_format = "yyyy-mm-dd hh:mm"
            if stripe:
                cell.fill = STRIPE_FILL

        # Variable groups
        row_issues = issues.get(id_key, {})
        for var_key in var_order:
            c0 = group_col_start[var_key]
            recs = row_issues.get(var_key)

            for offset in range(4):
                cell = ws.cell(row=r_ptr, column=c0 + offset)
                is_left_edge = (offset == 0)
                cell.border = Border(
                    left=THICK if is_left_edge else THIN, right=THIN, top=THIN, bottom=THIN,
                )
                cell.alignment = LEFT_WRAP

            if recs:
                orig_value = recs[0]["value"]
                reasons = "; ".join(dict.fromkeys(rec["issue"] for rec in recs))  # de-dup, keep order

                ws.cell(row=r_ptr, column=c0 + 0, value=("" if pd.isna(orig_value) else orig_value)).font = BODY_FONT
                ws.cell(row=r_ptr, column=c0 + 1, value=reasons).font = BODY_FONT

                for offset in (2, 3):
                    cell = ws.cell(row=r_ptr, column=c0 + offset)
                    cell.font = INPUT_FONT
                    cell.fill = INPUT_FILL
            else:
                if stripe:
                    for offset in range(4):
                        ws.cell(row=r_ptr, column=c0 + offset).fill = STRIPE_FILL

        r_ptr += 1

    # ── Column widths ───────────────────────────────────────────────────────
    for i, col_name in enumerate(id_cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ID_WIDTHS.get(col_name, 16)

    for var_key in var_order:
        c0 = group_col_start[var_key]
        widths = [16, 34, 16, 28]
        for offset, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(c0 + offset)].width = w

    # ── Freeze panes: keep headers + ID columns visible while scrolling ────
    freeze_col_letter = get_column_letter(n_id_cols + 1)
    ws.freeze_panes = f"{freeze_col_letter}{DATA_START_ROW}"

    # Filters on the header row for easy sorting/filtering by supervisors
    ws.auto_filter.ref = f"A{COL_HEADER_ROW}:{get_column_letter(n_cols)}{max(r_ptr - 1, COL_HEADER_ROW)}"

    return ws


def _id_cols_for(candidates, df):
    return [c for c in candidates if c in df.columns]


# ── Main ─────────────────────────────────────────────────────────────────────
def run_correction_template(data_path, batch_name=None, date_start=None, date_end=None, date_col="surveyDate", dict_path=None):
    compute_result = compute_quality_results_by_sheet(data_path, date_start, date_end, date_col, dict_path=dict_path)
    issues, var_order, var_labels = build_correction_maps(compute_result)

    batch = batch_name or os.path.basename(data_path).rsplit(".", 1)[0]
    main_sheet_name = compute_result["main_sheet_name"]
    id_col = compute_result["id_col"]

    wb = Workbook()
    wb.remove(wb.active)  # replaced by explicit per-sheet tabs below
    used_titles = set()

    sheet_summary = []

    # ── General/main sheet tab first — keeps the original tab name so the
    #    single-sheet/CSV workflow output is unchanged ──────────────────────
    main_report  = compute_result["sheets"][main_sheet_name]
    main_id_cols = _id_cols_for(ID_COLS, main_report)
    main_title   = _safe_sheet_title("Correction template", used_titles)
    _write_correction_sheet(
        wb, main_title, main_id_cols,
        issues.get(main_sheet_name, {}), var_order.get(main_sheet_name, []), var_labels.get(main_sheet_name, {}),
        main_report, id_col, batch, main_sheet_name,
    )
    sheet_summary.append((main_title, main_sheet_name, len(issues.get(main_sheet_name, {})), len(var_order.get(main_sheet_name, []))))

    # ── One tab per loop/repeat sheet ───────────────────────────────────────
    for sheet_name, report_df in compute_result["sheets"].items():
        if sheet_name == main_sheet_name:
            continue
        loop_id_field = compute_result["sheet_id_field"][sheet_name]
        loop_id_cols  = _id_cols_for(ID_COLS_LOOP, report_df)
        title = _safe_sheet_title(f"Loop - {sheet_name}", used_titles)
        _write_correction_sheet(
            wb, title, loop_id_cols,
            issues.get(sheet_name, {}), var_order.get(sheet_name, []), var_labels.get(sheet_name, {}),
            report_df, loop_id_field, batch, sheet_name,
        )
        sheet_summary.append((title, sheet_name, len(issues.get(sheet_name, {})), len(var_order.get(sheet_name, []))))

    os.makedirs(CORRECTIONS, exist_ok=True)
    out_path = os.path.join(CORRECTIONS, f"correction_template_{batch}.xlsx")
    wb.save(out_path)

    total_flagged = sum(n for _, _, n, _ in sheet_summary)
    print(f"✅  Correction template saved: {out_path}")
    for title, sheet_name, n_flag, n_vars in sheet_summary:
        print(f"    [{title}] (sheet '{sheet_name}'): {n_flag} flagged records, {n_vars} variable groups")
    print(f"    Total flagged records across all sheets: {total_flagged}")
    return out_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="AGEVAL Correction Template Generator")
    parser.add_argument("--data",  required=True, help="Path to collected data file (.xlsx with general + loop/repeat sheets, or legacy flat .csv)")
    parser.add_argument("--batch", default=None,  help="Batch name (optional label)")
    parser.add_argument("--date-start", default=None, help="Filter: only records on/after this date (YYYY-MM-DD)")
    parser.add_argument("--date-end",   default=None, help="Filter: only records on/before this date (YYYY-MM-DD)")
    parser.add_argument("--date-col",   default="surveyDate", help="Column name holding the survey date")
    parser.add_argument("--dict",       default=None, help="Path to personalized dictionary CSV (default: config path)")
    args = parser.parse_args()
    run_correction_template(args.data, args.batch, args.date_start, args.date_end, args.date_col, dict_path=args.dict)